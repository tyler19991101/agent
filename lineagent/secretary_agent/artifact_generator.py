import os
import uuid
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, List

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from secretary_agent.utils import sanitize_filename


@dataclass
class GeneratedArtifact:
    format: str
    filename: str
    path: str
    token: str
    url: str


class ArtifactGenerator:
    def __init__(self, output_dir: str, public_base_url: str = ""):
        self.output_dir = output_dir
        self.public_base_url = public_base_url.rstrip("/")
        os.makedirs(self.output_dir, exist_ok=True)

    def generate(
        self,
        *,
        title: str,
        content: str,
        output_formats: List[str],
    ) -> List[GeneratedArtifact]:
        generated: List[GeneratedArtifact] = []
        safe_title = sanitize_filename(title or "report", default="report")
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

        for fmt in output_formats:
            token = uuid.uuid4().hex
            filename = f"{safe_title}_{timestamp}.{fmt}"
            path = os.path.join(self.output_dir, filename)
            if fmt == "txt":
                self._write_txt(path, content)
            elif fmt == "docx":
                self._write_docx(path, title, content)
            elif fmt == "pdf":
                self._write_pdf(path, title, content)
            else:
                continue
            url = f"{self.public_base_url}/downloads/{token}" if self.public_base_url else ""
            generated.append(GeneratedArtifact(format=fmt, filename=filename, path=path, token=token, url=url))
        return generated

    @staticmethod
    def _write_txt(path: str, content: str) -> None:
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(content)

    @staticmethod
    def _write_docx(path: str, title: str, content: str) -> None:
        document = Document()
        if title:
            document.add_heading(title, level=1)
        for block in content.split("\n\n"):
            text = block.strip()
            if not text:
                continue
            document.add_paragraph(text)
        document.save(path)

    @staticmethod
    def _write_pdf(path: str, title: str, content: str) -> None:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        pdf = canvas.Canvas(path, pagesize=A4)
        pdf.setTitle(title or "report")
        pdf.setFont("STSong-Light", 16)
        width, height = A4
        y = height - 48
        if title:
            pdf.drawString(48, y, title)
            y -= 28
        pdf.setFont("STSong-Light", 11)
        for paragraph in content.split("\n"):
            line = paragraph.strip()
            if not line:
                y -= 14
            else:
                pdf.drawString(48, y, line[:90])
                y -= 16
            if y < 48:
                pdf.showPage()
                pdf.setFont("STSong-Light", 11)
                y = height - 48
        pdf.save()
